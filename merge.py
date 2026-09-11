import csv,json,os,sys
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from openpyxl.utils import get_column_letter

src=Path(sys.argv[1]); out=Path(sys.argv[2]); out.mkdir(parents=True,exist_ok=True)
expected=int(os.environ.get('EXPECTED_SHARDS','16'))

def read_all(name):
    rows=[]; fields=[]
    for p in src.rglob(name):
        with open(p,encoding='utf-8-sig',newline='') as f:
            r=csv.DictReader(f)
            if r.fieldnames and not fields: fields=r.fieldnames
            rows.extend(list(r))
    return fields,rows

img_fields,images=read_all('all_images.csv'); cov_fields,coverage=read_all('coverage.csv')
seen_shards=set()
for p in src.rglob('shard_info.json'):
    try: seen_shards.add(int(json.loads(p.read_text())['shard_index']))
    except: pass
missing=sorted(set(range(expected))-seen_shards)
seen=set(); ded=[]
for r in images:
    k=(r.get('category_url',''),r.get('category_page_url',''),r.get('product_key',''),r.get('image_url',''))
    if k in seen: continue
    seen.add(k); ded.append(r)
images=ded
fails=[r for r in images if r.get('status') in ('FAIL','ERROR')]
hard=[r for r in images if r.get('status')=='FAIL']; errs=[r for r in images if r.get('status')=='ERROR']
passes=[r for r in images if r.get('status') in ('PASS','PASS - LOW RES')]

def write(path,fields,rows):
    with open(path,'w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields or ['No data']); w.writeheader()
        if fields:
            for r in rows: w.writerow({k:r.get(k,'') for k in fields})
write(out/'all_images.csv',img_fields,images); write(out/'failed_images.csv',img_fields,fails); write(out/'coverage.csv',cov_fields,coverage)

by=defaultdict(lambda:{'p':set(),'i':0,'pass':0,'low':0,'fail':0,'err':0})
for r in images:
    k=(r.get('category_name',''),r.get('category_url','')); d=by[k]; d['i']+=1
    if r.get('product_key'): d['p'].add(r['product_key'])
    s=r.get('status'); d['pass']+=s=='PASS'; d['low']+=s=='PASS - LOW RES'; d['fail']+=s=='FAIL'; d['err']+=s=='ERROR'
sfields=['category_name','category_url','unique_products','images_checked','pass_1000_plus','pass_500_to_999','fail_below_minimum','errors','minimum_compliance_pct']
srows=[]
for (n,u),d in sorted(by.items()):
    den=max(1,d['i']-d['err']); comp=d['pass']+d['low']
    srows.append({'category_name':n,'category_url':u,'unique_products':len(d['p']),'images_checked':d['i'],'pass_1000_plus':d['pass'],'pass_500_to_999':d['low'],'fail_below_minimum':d['fail'],'errors':d['err'],'minimum_compliance_pct':round(comp*100/den,2)})
write(out/'category_summary.csv',sfields,srows)

wb=Workbook(); wb.remove(wb.active)
def sheet(name,headers,rows):
    ws=wb.create_sheet(name); ws.append(headers)
    for c in ws[1]: c.font=Font(bold=True)
    for r in rows: ws.append([r.get(h,'') for h in headers])
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for i,h in enumerate(headers,1):
        vals=[str(h)]+[str(r.get(h,'')) for r in rows[:150]]; ws.column_dimensions[get_column_letter(i)].width=min(55,max(10,max(map(len,vals))+2))
    return ws
comp=round(len(passes)*100/max(1,len(images)-len(errs)),2)
exec_rows=[{'Metric':'Expected crawl shards','Value':expected},{'Metric':'Completed crawl shards','Value':len(seen_shards)},{'Metric':'Missing crawl shards','Value':','.join(map(str,missing)) if missing else 'None'},{'Metric':'Category coverage rows','Value':len(coverage)},{'Metric':'Categories requiring review','Value':sum(1 for r in coverage if r.get('status')!='COMPLETE')},{'Metric':'Images checked','Value':len(images)},{'Metric':'Images below 500x500','Value':len(hard)},{'Metric':'Image errors','Value':len(errs)},{'Metric':'Minimum compliance %','Value':comp}]
sheet('Executive Summary',['Metric','Value'],exec_rows); cws=sheet('Coverage',cov_fields or ['No data'],coverage); fws=sheet('Failures',img_fields or ['No data'],fails); sheet('Category Summary',sfields,srows)
if img_fields: sheet('All Images',img_fields,images[:1000000])
red=PatternFill('solid',fgColor='F4CCCC'); yellow=PatternFill('solid',fgColor='FFF2CC')
if cov_fields and 'status' in cov_fields:
    j=cov_fields.index('status')
    for row in cws.iter_rows(min_row=2):
        v=str(row[j].value or '')
        if v!='COMPLETE':
            fill=red if v=='ERROR' or v.startswith('INCOMPLETE') else yellow
            for c in row: c.fill=fill
wb.save(out/'GemsNY_Image_Audit.xlsx')
summary={'expected_shards':expected,'completed_shards':sorted(seen_shards),'missing_shards':missing,'categories':len(coverage),'images':len(images),'failures':len(hard),'errors':len(errs),'compliance_pct':comp}
(out/'run_summary.json').write_text(json.dumps(summary,indent=2)); print(json.dumps(summary,indent=2))
